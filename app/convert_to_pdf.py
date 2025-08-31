import os
from pathlib import Path

from fpdf import FPDF
from docx import Document

def txt_to_pdf(txt_path, pdf_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=12)
    with open(txt_path, "r", encoding="utf-8") as f:
        for line in f:
            pdf.cell(0, 10, txt=line.strip(), ln=True)
    pdf.output(pdf_path)

def docx_to_pdf(docx_path, pdf_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_font("Arial", size=12)
    doc = Document(docx_path)
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            pdf.cell(0, 10, txt=text, ln=True)
    pdf.output(pdf_path)

import sys

def main():
    import argparse

    parser = argparse.ArgumentParser(description="Convert .txt and .docx files in a folder to PDF.")
    parser.add_argument("folder", nargs="?", default=".", help="Folder containing documents (default: current directory)")
    args = parser.parse_args()

    base_dir = Path(args.folder).resolve()
    pdf_dir = base_dir / "pdfs"
    pdf_dir.mkdir(exist_ok=True)

    print(f"[INFO] Scanning folder: {base_dir}")
    print(f"[INFO] Output PDFs will be saved in: {pdf_dir}")

    found = False
    converted = 0
    for file in base_dir.iterdir():
        if file.is_file() and file.suffix.lower() == ".txt":
            pdf_path = pdf_dir / (file.stem + ".pdf")
            print(f"[TXT]  {file.name} -> {pdf_path.name}")
            try:
                txt_to_pdf(file, pdf_path)
                print(f"   [OK] Converted {file.name}")
                converted += 1
            except Exception as e:
                print(f"   [ERR] Failed to convert {file.name}: {e}")
            found = True
        elif file.is_file() and file.suffix.lower() == ".docx":
            pdf_path = pdf_dir / (file.stem + ".pdf")
            print(f"[DOCX] {file.name} -> {pdf_path.name}")
            try:
                docx_to_pdf(file, pdf_path)
                print(f"   [OK] Converted {file.name}")
                converted += 1
            except Exception as e:
                print(f"   [ERR] Failed to convert {file.name}: {e}")
            found = True
    if not found:
        print("[INFO] No .txt or .docx files found in", base_dir)
    else:
        print(f"[DONE] Converted {converted} file(s) to PDF in {pdf_dir}")

if __name__ == "__main__":
    main()